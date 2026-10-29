#!/usr/bin/env python3
"""
Google Maps POI, Commercial & Non-Residential Multi-Threaded Parser
Parses businesses, shops, cafes, services, furniture, pickup points, offices, institutions,
and company buildings around given coordinates without using Google API.
Runs concurrent browser workers (multi-threading) for high-speed scraping.
Filters out pure residential houses where no commercial/institutional entities exist.
Supports multiple businesses per address, radius up to 10 km, tens of thousands of places,
and graceful early stop (Ctrl+C).
"""

import os
import re
import sys
import time
import math
import json
import random
import signal
import logging
import shutil
import queue
import threading
import urllib.parse
from typing import Optional, List, Tuple, Dict, Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    TimeoutException,
    NoSuchElementException,
    StaleElementReferenceException,
)

# ── Logging ──────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

# ── Constants ─────────────────────────────────────────────────────────────────
MAPS_SEARCH_URL = "https://www.google.com/maps/search/{query}/@{lat},{lon},{zoom}z"
MAPS_DIRECTIONS_URL = (
    "https://www.google.com/maps/dir/{from_lat},{from_lon}/{to_lat},{to_lon}"
    "/data=!4m2!4m1!3e2"
)
DIRECTION_TIMEOUT = 12

# ANSI colours
try:
    _USE_COLOR = sys.stdout.isatty()
except Exception:
    _USE_COLOR = False

def _c(code, text):
    return f"\033[{code}m{text}\033[0m" if _USE_COLOR else text

BOLD  = lambda t: _c("1", t)
DIM   = lambda t: _c("2", t)
CYAN  = lambda t: _c("96", t)
GREEN = lambda t: _c("92", t)
YELLOW= lambda t: _c("93", t)
RED   = lambda t: _c("91", t)
BLUE  = lambda t: _c("94", t)
GRAY  = lambda t: _c("37", t)

# ── Early stop signal handling ────────────────────────────────────────────────
_stop_requested = False

def _signal_handler(signum, frame):
    global _stop_requested
    if not _stop_requested:
        _stop_requested = True
        print(f"\n\n  {YELLOW('⚠')} Early stop requested by user (Ctrl+C).")
        print(f"  {CYAN('⟳')} Finishing current step and saving all collected data to Excel before exit...\n")
    else:
        print(f"\n  {RED('✗')} Force stopping immediately...")
        sys.exit(0)

signal.signal(signal.SIGINT, _signal_handler)
signal.signal(signal.SIGTERM, _signal_handler)

def _thread_excepthook(args):
    # Suppress tracebacks from daemon worker threads on early termination / interrupt
    if not _stop_requested:
        log.debug(f"Worker thread exception: {args.exc_value}")

threading.excepthook = _thread_excepthook

# ── Distance helpers ──────────────────────────────────────────────────────────
def haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    R = 6371.0
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    dphi    = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)
    a = math.sin(dphi/2)**2 + math.cos(phi1)*math.cos(phi2)*math.sin(dlambda/2)**2
    return R * 2 * math.asin(math.sqrt(a))

def fmt_km(km: float) -> str:
    if km < 1:
        return f"{int(km*1000)} m"
    return f"{km:.2f} km"

# ── Pure Residential Filter ───────────────────────────────────────────────────
RESIDENTIAL_CATEGORIES = {
    "жилой дом", "житловий будинок",
    "многоквартирный дом", "многоквартирный жилой дом", "багатоквартирний будинок",
    "жилой комплекс", "житловий комплекс",
    "apartment building", "residential building", "condominium complex", "apartment complex",
    "частный дом", "приватний будинок", "индивидуальный жилой дом",
    "дом", "здание", "building", "premise"
}

BUSINESS_KEYWORDS = [
    "магазин", "кафе", "ресторан", "бар", "салон", "парикмахерская",
    "аптека", "оптика", "стоматология", "клиника", "мед", "врач", "больница",
    "офис", "компания", "предприятие", "организация", "агентство",
    "пункт выдачи", "доставка", "мастерская", "ремонт", "сервис",
    "мебель", "авто", "мойка", "шиномонтаж", "банк", "банкомат",
    "клуб", "школа", "сад", "центр", "нотариус", "юрист", "бухгалтер",
    "гостиница", "отель", "хостел", "общежитие", "торговый", "склад",
    "производство", "завод", "фабрика", "цех", "кооператив", "фирма",
    "посуда", "цветы", "зоо", "одежда", "обувь", "техника",
    "церковь", "храм", "собор", "приход", "религия", "church",
    "спорт", "дюсш", "фитнес", "стадион", "зал", "секция", "sport",
    "увд", "полиция", "мвд", "росгвардия", "управление", "отдел", "police",
    "учреждение", "администрация", "суд", "мфц", "служба", "департамент",
    "скупка", "выкуп", "автосалон", "дилер", "авторынок",
    "колледж", "институт", "университет", "лицей", "гимназия"
]

def is_pure_residential(item: Dict[str, Any]) -> bool:
    """
    Returns True if the item is purely a residential house / apartment building
    without any business, shop, or commercial service inside.
    """
    cat = (item.get("category") or "").strip().lower()
    name = (item.get("name") or "").strip().lower()

    # If category or name contains any business/commercial keyword -> NOT pure residential (KEEP IT)
    for kw in BUSINESS_KEYWORDS:
        if kw in cat or kw in name:
            return False

    # If it is explicitly classified as a residential house / apartment building
    if cat in ("жилой дом", "житловий будинок", "многоквартирный дом", "многоквартирный жилой дом", "apartment building"):
        return True

    for res_cat in RESIDENTIAL_CATEGORIES:
        if res_cat == cat or cat.startswith(res_cat):
            if not item.get("phone") and not item.get("website"):
                return True

    # If name looks like a plain street address and no business name/category
    if re.match(r"^(ул\.|улица|пр\.|проспект|пер\.|переулок|шоссе|дом)\s+", name, re.I):
        if not item.get("phone") and not item.get("website") and not item.get("category"):
            return True

    return False

# ── UI helpers ────────────────────────────────────────────────────────────────
def _width():
    try:
        return shutil.get_terminal_size().columns
    except Exception:
        return 72

def _hr(char="─"):
    return char * min(_width(), 72)

def _banner():
    print()
    print(CYAN(_hr("═")))
    print(CYAN("  ") + BOLD("🗺  Google Maps Multi-Threaded Commercial & POI Parser"))
    print(CYAN("  ") + DIM("Multi-browser parallel scanning · Radius up to 10 km · Stores & Companies"))
    print(CYAN(_hr("═")))
    print()

def _section(title):
    print()
    print(YELLOW(f"  ▸ {title}"))
    print(GRAY("  " + _hr()))

def _readline() -> str:
    try:
        line = sys.stdin.buffer.readline()
        if not line:
            raise EOFError
        return line.decode("utf-8", errors="replace").rstrip("\n\r")
    except AttributeError:
        return input()

def _ask(prompt, default=None, validate=None, optional=False):
    hint = f" [{CYAN(str(default))}]" if default is not None else ""
    while True:
        try:
            sys.stdout.write(f"  {prompt}{hint}: ")
            sys.stdout.flush()
            raw = _readline().strip()
        except (EOFError, KeyboardInterrupt):
            print()
            sys.exit(0)
        if not raw and default is not None:
            raw = str(default)
        if not raw and not optional:
            print(RED("  ✗ Required. Please enter a value."))
            continue
        if validate:
            ok, msg = validate(raw)
            if not ok:
                print(RED(f"  ✗ {msg}"))
                continue
        return raw

def _ask_float(prompt, default=None, lo=None, hi=None):
    def v(s):
        try:
            val = float(s)
            if lo is not None and val < lo:
                return False, f"Must be at least {lo}"
            if hi is not None and val > hi:
                return False, f"Must be at most {hi}"
            return True, ""
        except ValueError:
            return False, "Please enter a valid decimal number (e.g. 2.5)"
    return float(_ask(prompt, default=str(default) if default is not None else None, validate=v))

def _ask_int(prompt, default=None, lo=1, hi=50000):
    def v(s):
        try:
            n = int(s)
            if not (lo <= n <= hi):
                return False, f"Must be between {lo} and {hi}"
            return True, ""
        except ValueError:
            return False, "Please enter a whole number"
    return int(_ask(prompt, default=str(default) if default is not None else None, validate=v))

def _ask_yn(prompt, default="y"):
    raw = _ask(prompt + " (y/n)", default=default).lower()
    return raw in ("y", "yes", "1", "т", "так", "д", "да")

def progress_bar(current, total, width=30):
    filled = int(width * current / total) if total else 0
    bar = "█" * filled + "░" * (width - filled)
    pct = int(100 * current / total) if total else 0
    return f"[{GREEN(bar)}] {pct:3d}%  {current}/{total}"

# ── Grid Generation ───────────────────────────────────────────────────────────
def generate_grid_points(lat: float, lon: float, radius_km: float) -> Tuple[List[Tuple[float, float]], float, float]:
    """
    Generate an overlapping geographic grid covering circle of radius_km.
    Step size (350m - 750m) ensures thorough coverage of all buildings.
    Points are sorted by distance from center outwards.
    Returns: (list_of_lat_lon, zoom_level, step_km)
    """
    if radius_km <= 1.0:
        step_km = 0.35  # ~350 m step - fine-grained street level
        zoom = 17.0
    elif radius_km <= 3.0:
        step_km = 0.45  # ~450 m step
        zoom = 17.0
    elif radius_km <= 6.0:
        step_km = 0.60  # ~600 m step
        zoom = 16.5
    else:
        step_km = 0.75  # ~750 m step (up to 10 km)
        zoom = 16.0

    d_lat = step_km / 111.32
    cos_lat = max(math.cos(math.radians(lat)), 0.05)
    d_lon = step_km / (111.32 * cos_lat)

    num_steps_lat = int(math.ceil(radius_km / step_km))
    num_steps_lon = int(math.ceil(radius_km / (step_km * cos_lat)))

    points_with_dist = []
    # Always include origin first
    points_with_dist.append((lat, lon, 0.0))

    for i in range(-num_steps_lat, num_steps_lat + 1):
        for j in range(-num_steps_lon, num_steps_lon + 1):
            if i == 0 and j == 0:
                continue
            p_lat = lat + i * d_lat
            p_lon = lon + j * d_lon
            dist = haversine_km(lat, lon, p_lat, p_lon)
            if dist <= radius_km:
                points_with_dist.append((p_lat, p_lon, dist))

    # Sort from origin outward (closest first)
    points_with_dist.sort(key=lambda p: p[2])
    coords = [(p[0], p[1]) for p in points_with_dist]
    return coords, zoom, step_km

# ── Interactive wizard & Non-interactive ENV support ──────────────────────────
def wizard() -> Dict[str, Any]:
    _banner()

    # Check if parameters were passed via Environment Variables (for non-interactive Docker runs)
    env_coords = os.environ.get("COORDS", "").strip()
    env_lat = os.environ.get("LAT", "").strip()
    env_lon = os.environ.get("LON", "").strip()

    if env_coords or (env_lat and env_lon):
        _section("Automated configuration detected from Environment Variables")
        if env_coords:
            parts = re.split(r"[,;\s]+", env_coords.replace(" ", ""))
            lat, lon = float(parts[0]), float(parts[1])
        else:
            lat, lon = float(env_lat), float(env_lon)

        radius_km = float(os.environ.get("RADIUS", "1.0"))
        scope_choice = int(os.environ.get("SCOPE", "1"))
        max_places = int(os.environ.get("MAX_PLACES", "1000"))
        num_threads = int(os.environ.get("THREADS", "3"))
        calc_walking = os.environ.get("WALKING", "n").lower() in ("y", "yes", "true", "1")
        excel_name = os.environ.get("EXCEL_NAME", "results")
        if excel_name.lower().endswith(".xlsx"):
            excel_name = excel_name[:-5]
        excel_name = excel_name + ".xlsx"

        filter_residential = True
        if scope_choice == 1:
            scope_name = "Вся коммерция, заведения и организации (без чистых жилых домов)"
            queries = [
                "магазин", "кафе", "услуги", "организации",
                "учреждения", "школа", "спорт", "авто", "церковь",
                "аптека", "пункт выдачи", "мебель"
            ]
            filter_residential = True
        elif scope_choice == 2:
            scope_name = "Только торговля и магазины"
            queries = ["магазин", "супермаркет", "мебель", "товары", "авто"]
            filter_residential = True
        elif scope_choice == 3:
            scope_name = "Еда и общепит"
            queries = ["кафе", "ресторан", "кофейня", "пекарня"]
            filter_residential = True
        elif scope_choice == 4:
            scope_name = "Все объекты включая чистые дома (без фильтрации)"
            queries = [
                "магазин", "кафе", "услуги", "организации",
                "учреждения", "школа", "спорт", "авто", "церковь", "дом"
            ]
            filter_residential = False
        else:
            custom_q = os.environ.get("CUSTOM_QUERY", "кафе, аптека")
            queries = [q.strip() for q in custom_q.split(",") if q.strip()]
            scope_name = f"Custom: {', '.join(queries)}"

        grid_pts, zoom, step_km = generate_grid_points(lat, lon, radius_km)

        print(f"  {BOLD('Origin:')}             {lat}, {lon}")
        print(f"  {BOLD('Radius:')}             {radius_km} km ({len(grid_pts)} grid cells, ~{int(step_km*1000)}m step, zoom {zoom})")
        print(f"  {BOLD('Scope:')}              {scope_name}")
        print(f"  {BOLD('Filter empty houses:')} {'yes' if filter_residential else 'no'}")
        print(f"  {BOLD('Categories:')}         {', '.join(queries)}")
        print(f"  {BOLD('Max places:')}         {max_places}")
        print(f"  {BOLD('Threads / Workers:')}  {GREEN(f'{num_threads} parallel browsers')}")
        print(f"  {BOLD('Walking route:')}      {'yes' if calc_walking else 'no'}")
        print(f"  {BOLD('Excel:')}              {excel_name}")
        print(CYAN(_hr()))
        print()

        return dict(
            lat=lat,
            lon=lon,
            radius_km=radius_km,
            scope_name=scope_name,
            queries=queries,
            filter_residential=filter_residential,
            max_places=max_places,
            num_threads=num_threads,
            calc_walking=calc_walking,
            excel_name=excel_name,
            grid_pts=grid_pts,
            zoom=zoom,
            step_km=step_km,
        )

    _section("Origin coordinates")
    def _parse_coords(s):
        s = s.strip().replace(" ", "")
        parts = re.split(r"[,;\s]+", s)
        if len(parts) == 2:
            try:
                la, lo = float(parts[0]), float(parts[1])
                if -90 <= la <= 90 and -180 <= lo <= 180:
                    return True, ""
                return False, "Latitude must be -90..90, longitude -180..180"
            except ValueError:
                pass
        return False, "Enter as: 64.56459,39.83801"

    raw_coords = _ask("Coordinates (lat,lon)", validate=_parse_coords)
    parts = re.split(r"[,;\s]+", raw_coords.strip().replace(" ", ""))
    lat, lon = float(parts[0]), float(parts[1])

    _section("Search radius")
    radius_km = _ask_float(
        "Search radius in km (0.1 - 10.0)",
        default=1.0,
        lo=0.1,
        hi=10.0,
    )

    _section("Scope of search (Режим скопа)")
    print(f"    {BOLD('1')}. {GREEN('Вся коммерция, заведения и организации')} [РЕКОМЕНДУЕТСЯ]")
    print(f"       {GRAY('Магазины, кафе, услуги, мебель, аптеки, ПВЗ, офисы в жилых домах и здания компаний.')}")
    print(f"       {YELLOW('Чистые жилые дома без заведений автоматически исключаются.')}")
    print(f"       {CYAN('На одном адресе сохраняются ВСЕ отдельные заведения.')}")
    print(f"    {BOLD('2')}. {CYAN('Только торговля и магазины')}")
    print(f"       {GRAY('Продукты, одежда, мебель, хозтовары, электроника.')}")
    print(f"    {BOLD('3')}. {CYAN('Еда и общепит')}")
    print(f"       {GRAY('Кафе, рестораны, кофейни, пекарни, бары, фастфуд.')}")
    print(f"    {BOLD('4')}. {GRAY('Все объекты включая чистые жилые дома (без фильтрации)')}")
    print(f"    {BOLD('5')}. {GRAY('Пользовательский запрос (ввести свои ключевые слова)')}")
    print()

    scope_choice = _ask_int("Select scope mode (1-5)", default=1, lo=1, hi=5)

    filter_residential = True

    if scope_choice == 1:
        scope_name = "Вся коммерция, заведения и организации (без чистых жилых домов)"
        queries = [
            "магазин", "кафе", "услуги", "организации",
            "учреждения", "школа", "спорт", "авто", "церковь",
            "аптека", "пункт выдачи", "мебель"
        ]
        filter_residential = True
    elif scope_choice == 2:
        scope_name = "Только торговля и магазины"
        queries = ["магазин", "супермаркет", "мебель", "товары", "авто"]
        filter_residential = True
    elif scope_choice == 3:
        scope_name = "Еда и общепит"
        queries = ["кафе", "ресторан", "кофейня", "пекарня"]
        filter_residential = True
    elif scope_choice == 4:
        scope_name = "Все объекты включая чистые дома (без фильтрации)"
        queries = [
            "магазин", "кафе", "услуги", "организации",
            "учреждения", "школа", "спорт", "авто", "церковь", "дом"
        ]
        filter_residential = False
    else:
        custom_input = _ask(
            "Enter custom query (comma-separated, e.g. 'кафе, аптека, банк')",
            default="кафе",
        )
        queries = [q.strip() for q in custom_input.split(",") if q.strip()]
        scope_name = f"Custom: {', '.join(queries)}"
        filter_residential = _ask_yn("Exclude pure residential buildings without businesses?", default="y")

    _section("Collection limits & Speed")
    max_places = _ask_int(
        "Max places to collect (1 - 50000)",
        default=1000 if radius_km > 2.0 else 300,
        lo=1,
        hi=50000,
    )

    num_threads = _ask_int(
        "Parallel browser threads for speedup (1 - 6)",
        default=3,
        lo=1,
        hi=6,
    )

    calc_walking = _ask_yn(
        "Calculate walking route distance? (slower, extra navigation per place)",
        default="y" if max_places <= 50 else "n",
    )

    _section("Output")
    excel_name = _ask("Excel filename (without .xlsx)", default="results")
    if excel_name.lower().endswith(".xlsx"):
        excel_name = excel_name[:-5]
    excel_name = excel_name + ".xlsx"

    grid_pts, zoom, step_km = generate_grid_points(lat, lon, radius_km)

    print()
    print(CYAN(_hr()))
    print(f"  {BOLD('Origin:')}             {lat}, {lon}")
    print(f"  {BOLD('Radius:')}             {radius_km} km ({len(grid_pts)} grid cells, ~{int(step_km*1000)}m step, zoom {zoom})")
    print(f"  {BOLD('Scope:')}              {scope_name}")
    print(f"  {BOLD('Filter empty houses:')} {'yes (только заведения и коммерция)' if filter_residential else 'no'}")
    print(f"  {BOLD('Categories:')}         {', '.join(queries)}")
    print(f"  {BOLD('Max places:')}         {max_places}")
    print(f"  {BOLD('Threads / Workers:')}  {GREEN(f'{num_threads} parallel browsers')}")
    print(f"  {BOLD('Walking route:')}      {'yes' if calc_walking else 'no (straight-line distance is always calculated)'}")
    print(f"  {BOLD('Excel:')}              {excel_name}")
    print(f"  {BOLD('Early stop:')}         {GREEN('Ctrl+C anytime to stop and save results immediately')}")
    print(CYAN(_hr()))
    print()

    return dict(
        lat=lat,
        lon=lon,
        radius_km=radius_km,
        scope_name=scope_name,
        queries=queries,
        filter_residential=filter_residential,
        max_places=max_places,
        num_threads=num_threads,
        calc_walking=calc_walking,
        excel_name=excel_name,
        grid_pts=grid_pts,
        zoom=zoom,
        step_km=step_km,
    )

# ── Browser factory ───────────────────────────────────────────────────────────
def make_driver(headless=True) -> webdriver.Chrome:
    options = Options()
    if headless:
        options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-gpu")
    options.add_argument("--disable-software-rasterizer")
    options.add_argument("--blink-settings=imagesEnabled=false")
    options.add_argument("--disable-extensions")
    options.add_argument("--disable-default-apps")
    options.add_argument("--disable-background-networking")
    options.add_argument("--disable-sync")
    options.add_argument("--disable-translate")
    options.add_argument("--no-first-run")
    options.add_argument("--disable-blink-features=AutomationControlled")
    options.add_argument("--disable-infobars")
    options.add_argument("--window-size=1280,800")
    options.add_argument("--lang=ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7")
    options.add_argument(
        "--user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    )
    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    options.add_experimental_option("useAutomationExtension", False)
    # Block images and heavy media downloads to save RAM and CPU
    options.add_experimental_option("prefs", {
        "profile.managed_default_content_settings.images": 2,
        "profile.default_content_setting_values.notifications": 2,
        "profile.default_content_setting_values.media_stream": 2,
    })

    chromedriver_bin = "/usr/bin/chromedriver"
    if os.path.exists(chromedriver_bin):
        service = Service(executable_path=chromedriver_bin)
    else:
        service = Service()

    driver = webdriver.Chrome(service=service, options=options)
    try:
        driver.execute_cdp_cmd(
            "Page.addScriptToEvaluateOnNewDocument",
            {"source": "Object.defineProperty(navigator,'webdriver',{get:()=>undefined})"},
        )
    except Exception:
        pass
    return driver

def dismiss_consent(driver):
    try:
        btn = WebDriverWait(driver, 4).until(
            EC.element_to_be_clickable((
                By.XPATH,
                "//button[contains(.,'Accept') or contains(.,'Прийняти') "
                "or contains(.,'Принять') or contains(.,'Agree') or contains(.,'Alle akzeptieren')]"
            ))
        )
        btn.click()
        time.sleep(1)
    except TimeoutException:
        pass

# ── Walking distance ──────────────────────────────────────────────────────────
def get_walking_distance(driver, from_lat, from_lon, to_lat, to_lon) -> Optional[str]:
    url = MAPS_DIRECTIONS_URL.format(
        from_lat=from_lat, from_lon=from_lon,
        to_lat=to_lat,     to_lon=to_lon,
    )
    try:
        driver.get(url)
        dismiss_consent(driver)
        WebDriverWait(driver, DIRECTION_TIMEOUT).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "[data-trip-index]"))
        )
        time.sleep(1.2)
        for sel in [
            "[data-trip-index='0'] .fontBodyMedium",
            "[data-trip-index='0'] .fontBodyLarge",
            ".section-trip-summary",
        ]:
            try:
                for el in driver.find_elements(By.CSS_SELECTOR, sel):
                    text = el.text.strip()
                    if text and any(x in text for x in ["km", "m", "min", "хв", "мин", "км", "м"]):
                        return text
            except Exception:
                pass
        try:
            card_text = driver.find_element(By.CSS_SELECTOR, "[data-trip-index='0']").text
            m = re.search(r"[\d,\.]+\s*(км|м|km|m)", card_text)
            if m:
                return m.group(0)
        except Exception:
            pass
    except Exception as e:
        log.debug(f"Walking distance error: {e}")
    return None

# ── Fast Card Parser ──────────────────────────────────────────────────────────
def parse_card(card, origin_lat: float, origin_lon: float) -> Optional[Dict[str, Any]]:
    """
    Extracts all fields directly from search feed card (div.Nv2PK).
    Zero additional page loads required.
    """
    item: Dict[str, Any] = {}
    try:
        link = card.find_element(By.CSS_SELECTOR, "a.hfpxzc, a[href*='/maps/place/']")
        href = link.get_attribute("href") or ""
        item["url"] = href.split("?")[0]
        item["name"] = link.get_attribute("aria-label") or ""

        # Extract coordinates from URL: !3d<lat>!4d<lon>
        m = re.search(r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)", href)
        if m:
            item["lat"], item["lon"] = float(m.group(1)), float(m.group(2))
        else:
            m2 = re.search(r"@(-?\d+\.\d+),(-?\d+\.\d+)", href)
            if m2:
                item["lat"], item["lon"] = float(m2.group(1)), float(m2.group(2))
    except Exception:
        return None

    if not item.get("name"):
        try:
            name_el = card.find_element(By.CSS_SELECTOR, ".qBF1Pd")
            item["name"] = name_el.text.strip()
        except NoSuchElementException:
            pass

    # Rating
    try:
        r_el = card.find_element(By.CSS_SELECTOR, "span.MW4etd")
        r_text = r_el.text.strip()
        if r_text and re.match(r"^\d[,\.]\d$", r_text):
            item["rating"] = r_text
    except Exception:
        pass

    # Website
    for w in card.find_elements(
        By.CSS_SELECTOR,
        "a[aria-label*='сайт' i], a[aria-label*='website' i], a[data-value*='Сайт'], a[data-value*='Website']"
    ):
        w_href = w.get_attribute("href")
        if w_href and "google.com" not in w_href:
            item["website"] = w_href
            break

    # Parse text lines inside the card for category, address, phone, working hours
    try:
        raw_text = card.text or ""
    except Exception:
        raw_text = ""
    lines = [ln.strip() for ln in raw_text.split("\n") if ln.strip()]

    for line in lines:
        # Skip rating lines (e.g. "3,3" or "4,8(18)" or "4.5 (120)")
        if re.match(r"^[\d,\.]+\s*(\(\d+\))?$", line):
            continue

        # Check phone (+7 ..., 8 ..., etc.)
        if not item.get("phone"):
            ph = re.search(r"(\+?\d[\d\s\-\(\)]{7,}\d)", line)
            if ph and sum(c.isdigit() for c in ph.group(1)) >= 7:
                item["phone"] = ph.group(1).strip()
                continue

        # Check working hours
        if any(w in line.lower() for w in [
            "відчинено", "зачинено", "открыто", "закрыто",
            "цілодобово", "круглосуточно", "open", "closed", "відкриється", "откроется"
        ]):
            if not item.get("hours"):
                item["hours"] = line
                continue

        # Check category & address separated by '·'
        if "·" in line:
            parts = [p.strip() for p in line.split("·") if p.strip()]
            for p in parts:
                if re.search(r"(ул\.|улица|пр\.|проспект|пер\.|переулок|шоссе|дом|\d+)", p, re.I):
                    item.setdefault("address", p)
                elif not item.get("category") and not re.match(r"^[\d,\.]", p) and len(p) > 2:
                    if not any(ign in p.lower() for ign in ["маршрут", "сайт", "фото"]):
                        item["category"] = p
        elif re.search(r"(ул\.|улица|пр\.|проспект|пер\.|переулок|шоссе|дом\s*\d+)", line, re.I):
            item.setdefault("address", line)
        elif not item.get("category") and line != item.get("name") and not re.match(r"^[\d,\.]", line) and len(line) > 3:
            if not any(ign in line.lower() for ign in ["маршрут", "фото", "відгук", "отзыв", "сайт"]):
                item.setdefault("category", line)

    # Compute straight-line distance
    if item.get("lat") is not None and item.get("lon") is not None:
        straight = haversine_km(origin_lat, origin_lon, item["lat"], item["lon"])
        item["distance_meters"] = int(round(straight * 1000))
        item["distance_straight"] = fmt_km(straight)
        item["distance_straight_km"] = round(straight, 4)

    return item

# ── Print single place ────────────────────────────────────────────────────────
def _print_place(d: Dict[str, Any], idx: int, total: int):
    print()
    print(f"  {progress_bar(idx, total)}")
    print(f"  {BOLD(GREEN(d.get('name', '(no name)')))}  {GRAY(d.get('category',''))}")
    rows = [
        ("📍 Address",   "address"),
        ("📞 Phone",     "phone"),
        ("🌐 Website",   "website"),
        ("🕐 Hours",     "hours"),
        ("⭐ Rating",    "rating"),
        ("📏 Straight",  "distance_straight"),
        ("🚶 Walking",   "distance_walking"),
    ]
    for label, key in rows:
        val = d.get(key)
        if val and val != "—":
            print(f"  {GRAY(label):<28} {val}")
    if d.get("lat"):
        print(f"  {GRAY('🔵 Coords'):<28} {d['lat']}, {d['lon']}")
    print(GRAY("  " + _hr()))

# ── Excel Export ──────────────────────────────────────────────────────────────
COLUMNS = [
    ("No", "idx", 6),
    ("Name", "name", 32),
    ("Address", "address", 38),
    ("Phone", "phone", 20),
    ("Website", "website", 35),
    ("Hours", "hours", 30),
    ("Rating", "rating", 10),
    ("Category", "category", 24),
    ("Latitude", "lat", 14),
    ("Longitude", "lon", 14),
    ("Straight Distance (m)", "distance_meters", 22),
    ("Walking Route", "distance_walking", 20),
    ("URL", "url", 50),
]

def export_excel(results: List[Dict[str, Any]], filepath: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Google Maps Places"

    header_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
    header_font = Font(name="Calibri", bold=True, color="E0E0E0", size=11)
    header_align= Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin   = Side(style="thin", color="444466")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_odd  = PatternFill(start_color="F5F5FF", end_color="F5F5FF", fill_type="solid")
    f_even = PatternFill(start_color="EAEAF7", end_color="EAEAF7", fill_type="solid")

    ws.row_dimensions[1].height = 36
    for col_i, (header, _, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_i, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
        ws.column_dimensions[get_column_letter(col_i)].width = width

    body_font = Font(name="Calibri", size=10)
    body_align= Alignment(vertical="center", wrap_text=True)
    body_align_r = Alignment(horizontal="right", vertical="center", wrap_text=True)
    body_align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
    link_font = Font(name="Calibri", size=10, color="4169E1", underline="single")

    # Sort results by proximity from origin (ascending meters) so closest places are first
    sorted_items = sorted(
        results,
        key=lambda x: (x.get("distance_meters") is None, x.get("distance_meters", 0))
    )

    for row_i, item in enumerate(sorted_items, 2):
        ws.row_dimensions[row_i].height = 22
        fill = f_odd if row_i % 2 else f_even
        item["idx"] = row_i - 1

        for col_i, (_, key, _) in enumerate(COLUMNS, 1):
            val = item.get(key, "")
            cell = ws.cell(row=row_i, column=col_i)
            cell.border = border
            cell.fill = fill
            cell.font = body_font
            cell.alignment = body_align

            if key == "idx":
                cell.value = row_i - 1
                cell.alignment = body_align_c
            elif key == "distance_meters":
                if isinstance(val, (int, float)):
                    cell.value = int(round(val))
                    cell.number_format = '#,##0 "m"'
                    cell.alignment = body_align_r
                else:
                    cell.value = val or ""
            elif key in ("lat", "lon") and isinstance(val, (int, float)):
                cell.value = float(val)
                cell.number_format = '0.000000'
                cell.alignment = body_align_r
            elif key == "rating" and val:
                try:
                    num_r = float(str(val).replace(",", "."))
                    cell.value = num_r
                    cell.number_format = '0.0'
                    cell.alignment = body_align_c
                except (ValueError, TypeError):
                    cell.value = str(val)
            elif key in ("url", "website") and val:
                cell.value = str(val)
                cell.hyperlink = str(val)
                cell.font = link_font
            else:
                cell.value = str(val) if val is not None else ""

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(filepath)

# ── Multi-Threaded Grid Worker ────────────────────────────────────────────────
def _worker_loop(
    worker_id: int,
    cell_queue: queue.Queue,
    total_cells: int,
    origin_lat: float,
    origin_lon: float,
    radius_km: float,
    queries: List[str],
    filter_residential: bool,
    max_places: int,
    calc_walking: bool,
    zoom: float,
    excel_path: str,
    lock: threading.Lock,
    print_lock: threading.Lock,
    seen_urls: set,
    results: list,
    last_saved: list,
):
    global _stop_requested
    driver = None
    try:
        driver = make_driver(headless=True)
    except Exception:
        return

    try:
        while not _stop_requested:
            with lock:
                if len(results) >= max_places:
                    break

            try:
                cell_info = cell_queue.get_nowait()
            except queue.Empty:
                break

            cell_idx, c_lat, c_lon = cell_info
            cell_dist = haversine_km(origin_lat, origin_lon, c_lat, c_lon)

            with print_lock:
                sys.stdout.write(
                    f"\r  {CYAN('⟳')} [Worker {worker_id} · Cell {cell_idx}/{total_cells} · {cell_dist:.2f}km] "
                    f"Collected {BOLD(str(len(results)))}/{max_places} places   "
                )
                sys.stdout.flush()

            try:
                for q in queries:
                    if _stop_requested:
                        break
                    with lock:
                        if len(results) >= max_places:
                            break

                    encoded_q = urllib.parse.quote(q)
                    search_url = MAPS_SEARCH_URL.format(query=encoded_q, lat=c_lat, lon=c_lon, zoom=zoom)

                    try:
                        driver.get(search_url)
                        dismiss_consent(driver)
                        time.sleep(1.8)
                    except Exception:
                        continue

                    scroll_el = None
                    for sel in ["div[role='feed']", ".m6QErb[aria-label]", ".m6QErb"]:
                        try:
                            scroll_el = driver.find_element(By.CSS_SELECTOR, sel)
                            break
                        except Exception:
                            pass

                    consecutive_empty = 0

                    for _ in range(5):
                        if _stop_requested:
                            break
                        with lock:
                            if len(results) >= max_places:
                                break

                        try:
                            cards = driver.find_elements(By.CSS_SELECTOR, "div.Nv2PK")
                        except Exception:
                            break

                        found_new = False

                        for card in cards:
                            if _stop_requested:
                                break
                            with lock:
                                if len(results) >= max_places:
                                    break

                            try:
                                link_el = card.find_element(By.CSS_SELECTOR, "a.hfpxzc, a[href*='/maps/place/']")
                                raw_href = (link_el.get_attribute("href") or "").split("?")[0]
                                if not raw_href:
                                    continue

                                with lock:
                                    if raw_href in seen_urls:
                                        continue
                                    seen_urls.add(raw_href)

                                item = parse_card(card, origin_lat, origin_lon)
                                if not item or not item.get("name"):
                                    continue

                                # Filter out pure residential houses without commercial/institutional presence
                                if filter_residential and is_pure_residential(item):
                                    continue

                                # Check radius tolerance
                                p_lat, p_lon = item.get("lat"), item.get("lon")
                                if p_lat is not None and p_lon is not None:
                                    d_from_origin = haversine_km(origin_lat, origin_lon, p_lat, p_lon)
                                    if d_from_origin > radius_km * 1.15:
                                        continue

                                if calc_walking and p_lat is not None and p_lon is not None:
                                    w_dist = get_walking_distance(driver, origin_lat, origin_lon, p_lat, p_lon)
                                    item["distance_walking"] = w_dist or "—"
                                else:
                                    item["distance_walking"] = "—"

                                with lock:
                                    results.append(item)
                                    count = len(results)

                                    # Periodic auto-save every 25 items
                                    if count - last_saved[0] >= 25:
                                        try:
                                            export_excel(list(results), excel_path)
                                            last_saved[0] = count
                                        except Exception:
                                            pass

                                found_new = True

                                with print_lock:
                                    _print_place(item, count, max_places)

                                with lock:
                                    if len(results) >= max_places:
                                        break
                            except Exception:
                                continue

                        if _stop_requested:
                            break
                        with lock:
                            if len(results) >= max_places:
                                break

                        if not found_new:
                            consecutive_empty += 1
                            if consecutive_empty >= 2:
                                break
                        else:
                            consecutive_empty = 0

                        try:
                            if scroll_el:
                                driver.execute_script("arguments[0].scrollTop += 1400;", scroll_el)
                            else:
                                driver.execute_script("window.scrollBy(0, 1400);")
                        except Exception:
                            pass
                        time.sleep(1.0)

                        try:
                            driver.find_element(By.CSS_SELECTOR, "span.HlvSq")
                            break
                        except Exception:
                            pass
            except Exception:
                pass
            finally:
                try:
                    cell_queue.task_done()
                except Exception:
                    pass
    except Exception:
        pass
    finally:
        if driver:
            try:
                driver.quit()
            except Exception:
                pass

# ── Multi-Threaded Grid Engine ────────────────────────────────────────────────
def collect_places_grid(
    origin_lat: float,
    origin_lon: float,
    radius_km: float,
    queries: List[str],
    filter_residential: bool,
    max_places: int,
    num_threads: int,
    calc_walking: bool,
    grid_pts: List[Tuple[float, float]],
    zoom: float,
    excel_name: str,
) -> List[Dict[str, Any]]:
    """
    Coordinates multiple parallel browser threads to scan grid cells concurrently.
    """
    global _stop_requested
    seen_urls = set()
    results: List[Dict[str, Any]] = []
    lock = threading.Lock()
    print_lock = threading.Lock()
    last_saved = [0]

    out_dir = "/output"
    os.makedirs(out_dir, exist_ok=True)
    excel_path = os.path.join(out_dir, excel_name)

    total_cells = len(grid_pts)
    _section(
        f"Starting Multi-Threaded Scan ({num_threads} browser threads · "
        f"{total_cells} cells · radius {radius_km} km · target {max_places} places)"
    )
    print(f"  {GRAY('Press Ctrl+C at any time to finish and save results immediately.')}\n")

    cell_queue = queue.Queue()
    for idx, (lat, lon) in enumerate(grid_pts, 1):
        cell_queue.put((idx, lat, lon))

    threads = []
    for w_id in range(1, num_threads + 1):
        t = threading.Thread(
            target=_worker_loop,
            kwargs=dict(
                worker_id=w_id,
                cell_queue=cell_queue,
                total_cells=total_cells,
                origin_lat=origin_lat,
                origin_lon=origin_lon,
                radius_km=radius_km,
                queries=queries,
                filter_residential=filter_residential,
                max_places=max_places,
                calc_walking=calc_walking,
                zoom=zoom,
                excel_path=excel_path,
                lock=lock,
                print_lock=print_lock,
                seen_urls=seen_urls,
                results=results,
                last_saved=last_saved,
            ),
            daemon=True,
        )
        threads.append(t)
        t.start()

    # Wait for all threads to finish or early stop
    try:
        while any(t.is_alive() for t in threads):
            if _stop_requested:
                break
            with lock:
                if len(results) >= max_places:
                    break
            time.sleep(0.5)
    except KeyboardInterrupt:
        _stop_requested = True
        print(f"\n\n  {YELLOW('⚠')} Interrupted by user. Waiting for threads to save and finish...")

    # Wait for threads to close browsers
    for t in threads:
        t.join(timeout=3.0)

    print()
    return results

# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    cfg = wizard()
    results = []

    try:
        results = collect_places_grid(
            origin_lat=cfg["lat"],
            origin_lon=cfg["lon"],
            radius_km=cfg["radius_km"],
            queries=cfg["queries"],
            filter_residential=cfg["filter_residential"],
            max_places=cfg["max_places"],
            num_threads=cfg["num_threads"],
            calc_walking=cfg["calc_walking"],
            grid_pts=cfg["grid_pts"],
            zoom=cfg["zoom"],
            excel_name=cfg["excel_name"],
        )
    except KeyboardInterrupt:
        print(f"\n\n  {YELLOW('⚠')} Process stopped by user (Ctrl+C).")

    # ── Final Summary & Save
    excel_name = cfg["excel_name"]
    out_dir = "/output"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, excel_name)

    print()
    print(CYAN(_hr("═")))
    if _stop_requested:
        print(f"  {BOLD(YELLOW('Scan stopped early by user!'))}  Saved {BOLD(str(len(results)))} places/businesses.")
    else:
        print(f"  {BOLD(GREEN('Done!'))}  Successfully collected {BOLD(str(len(results)))} places/businesses.")
    print(CYAN(_hr("═")))

    if results:
        export_excel(results, out_path)
        print(f"\n  {GREEN('✓')} Excel saved: {BOLD(out_path)} ({len(results)} rows)")
    else:
        print(f"\n  {YELLOW('!')} No places were collected.")

    # ── Sample JSON (or full if small)
    print()
    if len(results) <= 15 and results:
        print(json.dumps(results, ensure_ascii=False, indent=2))
    elif results:
        print(f"  {GRAY('Sample of first 3 items:')}")
        print(json.dumps(results[:3], ensure_ascii=False, indent=2))
        print(f"  {GRAY(f'... and {len(results)-3} more items in Excel: {excel_name}')}")

if __name__ == "__main__":
    main()